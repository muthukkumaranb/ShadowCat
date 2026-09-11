from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).parent
ARTIFACTS = ROOT / "artifacts" / "experiments" / "ucs_pca_20260904"
OUTPUT = ROOT / "artifacts" / "LSTM_UCS_Model_Report.docx"


def load_json(path: Path) -> dict:
    with path.open(encoding="utf-8") as file:
        return json.load(file)


def shade(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def set_cell_text(cell, text: str, bold: bool = False, color: str | None = None) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(str(text))
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(document: Document, headers: list[str], rows: list[list[str]], widths: list[float] | None = None):
    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for index, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[index], header, bold=True, color="FFFFFF")
        shade(table.rows[0].cells[index], "17324D")
    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            set_cell_text(cells[index], value)
            if len(table.rows) % 2 == 0:
                shade(cells[index], "EAF1F5")
    if widths:
        for row in table.rows:
            for index, width in enumerate(widths):
                row.cells[index].width = Inches(width)
    document.add_paragraph()
    return table


def add_bullet(document: Document, text: str) -> None:
    document.add_paragraph(text, style="List Bullet")


def metric(value: float) -> str:
    return f"{value:.4f}"


def main() -> None:
    comparison = load_json(ARTIFACTS / "comparison.json")
    best = load_json(ARTIFACTS / "pca_32" / "lstm_metrics.json")
    metadata = load_json(ARTIFACTS / "pca_32" / "lstm_ucs_experiment.json")

    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    normal = document.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.5)
    for style_name, size, color in [("Title", 28, "17324D"), ("Heading 1", 17, "17324D"), ("Heading 2", 12.5, "287B8E")]:
        style = document.styles[style_name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("LSTM + Unified Cyber State")
    run.bold = True
    run.font.size = Pt(28)
    run.font.color.rgb = RGBColor.from_string("17324D")
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("CSE-CIC-IDS2018 future attack prediction report")
    run.font.size = Pt(14)
    run.font.color.rgb = RGBColor.from_string("287B8E")
    date = document.add_paragraph()
    date.alignment = WD_ALIGN_PARAGRAPH.CENTER
    date.add_run("Experiment date: 4 September 2026 | Best documented PCA run: 32 components")
    document.add_paragraph()

    document.add_heading("Executive Summary", level=1)
    document.add_paragraph(
        "This project implements a temporal binary classifier that consumes canonical Unified Cyber State (UCS) windows derived from the CSE-CIC-IDS2018 network-flow dataset. The model observes 30 consecutive one-minute windows and predicts whether an attack will occur at the forecast target five windows ahead. The pipeline uses chronological splitting, purge and embargo protection, training-only normalization, and a causal PyTorch LSTM."
    )
    document.add_paragraph(
        f"Among the comparable PCA experiments, the 32-component model is the strongest saved result on the held-out test set: accuracy {metric(best['accuracy'])}, precision {metric(best['precision'])}, recall {metric(best['recall'])}, F1 {metric(best['f1'])}, ROC-AUC {metric(best['roc_auc'])}, and PR-AUC {metric(best['pr_auc'])}. Its classification threshold was selected on validation data ({best['threshold']:.2f}), rather than fixed blindly at 0.50."
    )

    document.add_heading("1. Problem and Data Flow", level=1)
    document.add_paragraph("The system is designed for early warning from temporally ordered network activity:")
    add_table(document, ["Stage", "What happens"], [
        ["Raw input", "CSE-CIC-IDS2018 CSV flow records"],
        ["Canonical representation", "UCS ingestion, cleaning, attack mapping, and one-minute aggregation"],
        ["Target", "future_attack_label over a five-window forecast horizon"],
        ["Temporal protection", "Chronological 70/15/15 split with a 35-window purge and embargo"],
        ["Preprocessing", "Training-only log1p for traffic-volume families, followed by robust scaling"],
        ["Sequence construction", "30 contiguous windows per sample, shaped as (30, features)"],
        ["Prediction", "LSTM probability of the future attack class"],
    ], [1.55, 5.7])

    document.add_heading("2. Model Architecture and Working", level=1)
    document.add_paragraph("The classifier is intentionally compact and causal. It processes the sequence in time order and uses only information available within the observed lookback window.")
    add_table(document, ["Component", "Configuration / role"], [
        ["Input", "30 time steps x 32 PCA features in the selected run"],
        ["Recurrent layer", "One unidirectional LSTM layer with hidden size 64"],
        ["Regularization", "Dropout 0.20 applied to the final hidden representation"],
        ["Output head", "Linear layer from 64 hidden units to one logit"],
        ["Probability", "Sigmoid converts the logit to a positive-class probability"],
        ["Decision", "Probability >= validation-selected threshold becomes an attack prediction"],
    ], [1.55, 5.7])
    document.add_paragraph("At each time step, the LSTM updates its hidden state using the current UCS feature vector and the previous state. After the thirtieth window, the final output summarizes the recent temporal context. Dropout regularizes that representation, the fully connected layer produces a logit, and the sigmoid converts it into an interpretable probability.")

    document.add_heading("3. Training and Evaluation Protocol", level=1)
    add_table(document, ["Setting", "Value"], [
        ["Loss", "BCEWithLogitsLoss with class weighting enabled in the authoritative configuration"],
        ["Optimizer", "Adam"],
        ["Learning rate", "0.001"],
        ["Weight decay", "0.0001"],
        ["Batch size", "64"],
        ["Maximum epochs", "30"],
        ["Early stopping", "Enabled; patience 5, minimum improvement 0.001"],
        ["Random seed", "42"],
        ["Model selection", "Validation PR-AUC, then validation F1"],
    ], [1.55, 5.7])
    document.add_paragraph("The purge and embargo width is 35 windows: 30 lookback windows plus 5 forecast-horizon windows. This reduces temporal leakage around split boundaries. Normalization and PCA are fitted using purged training windows only, then applied to validation and test data.")

    document.add_heading("4. Comparative Performance", level=1)
    rows = []
    for name, values in comparison["experiments"].items():
        rows.append([
            name,
            str(values["features"]),
            str(values["train_sequences"]),
            str(values["validation_sequences"]),
            str(values["test_sequences"]),
            metric(values["test_f1"]),
            metric(values["test_pr_auc"]),
        ])
    add_table(document, ["Run", "Features", "Train", "Val", "Test", "Test F1", "Test PR-AUC"], rows, [1.45, 0.7, 0.7, 0.55, 0.55, 0.75, 0.9])
    document.add_paragraph("The comparison uses the same sequence counts across all PCA settings. PCA-32 provides the best test PR-AUC and test F1 among these runs, despite using only 32 of the original 406 UCS features. The validation winner was selected by validation PR-AUC followed by validation F1, so test performance is reported as a final held-out assessment rather than a tuning target.")

    document.add_heading("5. Best Run: PCA-32 Detailed Results", level=1)
    add_table(document, ["Metric", "Result", "Meaning"], [
        ["Accuracy", metric(best["accuracy"]), "Correct predictions across both classes"],
        ["Precision", metric(best["precision"]), "Share of predicted attacks that were attacks"],
        ["Recall", metric(best["recall"]), "Share of attacks detected"],
        ["F1 score", metric(best["f1"]), "Balance between precision and recall"],
        ["ROC-AUC", metric(best["roc_auc"]), "Ranking quality across thresholds"],
        ["PR-AUC", metric(best["pr_auc"]), "Precision-recall ranking quality"],
        ["Decision threshold", f"{best['threshold']:.2f}", "Validation-selected probability cutoff"],
    ], [1.55, 1.05, 4.65])
    document.add_heading("Confusion Matrix", level=2)
    add_table(document, ["", "Predicted normal", "Predicted attack"], [
        ["Actual normal", str(best["tn"]), str(best["fp"])],
        ["Actual attack", str(best["fn"]), str(best["tp"])],
    ], [1.7, 2.6, 2.6])
    document.add_paragraph("The model detected 86 of 171 attack-labelled test sequences (recall 0.5029) and produced 42 false alarms. It missed 85 attacks, so the model is useful as an experimental early-warning baseline but is not yet sufficient for unattended security response.")

    document.add_heading("6. PCA Findings", level=1)
    document.add_paragraph(
        f"The selected PCA-32 transform was fitted on 406 input features. The first 32 components preserve {metadata['pca']['cumulative_explained_variance'][-1] * 100:.6f}% cumulative variance. The first component alone explains {metadata['pca']['explained_variance_ratio'][0] * 100:.4f}%, indicating strong redundancy or scale concentration in the original UCS feature space."
    )
    document.add_paragraph("Variance preservation does not guarantee equal predictive usefulness. In this experiment, PCA-32 outperformed PCA-64 and PCA-128 on the held-out test set, showing that the lower-dimensional representation was a better-performing operating point for this LSTM run.")

    document.add_heading("7. Strengths, Limitations, and Next Steps", level=1)
    document.add_heading("Strengths", level=2)
    add_bullet(document, "Uses a canonical UCS representation rather than feeding raw flow records directly into the LSTM.")
    add_bullet(document, "Preserves temporal order and uses a causal unidirectional model.")
    add_bullet(document, "Controls common temporal leakage paths with purge, embargo, contiguous sequences, and training-only transforms.")
    add_bullet(document, "Stores configuration, checkpoints, predictions, scaler/PCA metadata, and metrics as experiment artifacts.")
    document.add_heading("Limitations", level=2)
    add_bullet(document, "The best test F1 of 0.5753 and recall of 0.5029 leave substantial missed-attack risk.")
    add_bullet(document, "The report is based on saved experiment artifacts; operational latency, calibration, and performance under distribution shift were not evaluated here.")
    add_bullet(document, "PCA improves compactness but reduces direct feature interpretability.")
    add_bullet(document, "The comparison contains one saved run per PCA setting, so repeated-seed confidence intervals are not available.")
    document.add_heading("Recommended next steps", level=2)
    add_bullet(document, "Repeat the PCA-32 experiment across multiple seeds and report mean, standard deviation, and confidence intervals.")
    add_bullet(document, "Tune the operating threshold against an explicit security objective, such as minimum recall at an acceptable false-alarm rate.")
    add_bullet(document, "Add calibration and time-to-detection analysis before using probabilities operationally.")
    add_bullet(document, "Compare against a persistence/rule baseline and a non-recurrent model to quantify the value of temporal recurrence.")

    document.add_heading("8. Reproducibility References", level=1)
    document.add_paragraph("The implementation and artifacts used for this report are in the project workspace:")
    add_table(document, ["Item", "Path"], [
        ["Pipeline overview", "README.md"],
        ["Authoritative settings", "configs/model_config.yaml"],
        ["UCS preparation and PCA", "lstm/ucs.py"],
        ["LSTM architecture", "lstm/model.py"],
        ["Training loop", "lstm/trainer.py"],
        ["Evaluation metrics", "lstm/evaluator.py"],
        ["Best-run metrics", "artifacts/experiments/ucs_pca_20260904/pca_32/lstm_metrics.json"],
        ["PCA comparison", "artifacts/experiments/ucs_pca_20260904/comparison.json"],
        ["Best-run metadata", "artifacts/experiments/ucs_pca_20260904/pca_32/lstm_ucs_experiment.json"],
    ], [1.8, 5.45])
    document.add_paragraph("Generated from the saved artifacts on 4 September 2026.")

    for paragraph in document.paragraphs:
        paragraph.paragraph_format.space_after = Pt(5)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()